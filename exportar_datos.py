#!/usr/bin/env python3
"""
exportar_datos.py — Exporta las tablas de Pishico a CSV y Excel (respaldo).

Uso:
    cd ~/chatbot-shi
    source venv/bin/activate
    python3 exportar_datos.py

Genera una carpeta export_YYYYMMDD_HHMM/ con:
    - usuarios.csv
    - sesiones.csv
    - progreso_vocabulario.csv
    - progreso_cuento.csv
    - pishico_respaldo.xlsx   (todas las tablas, una por hoja)
    - resumen.txt             (conteos y rango de fechas)

Lee DATABASE_URL del .env igual que db.py, así apunta a la misma base
(PostgreSQL en RDS o SQLite local) sin configuración adicional.

Opcional — anonimizar antes de compartir:
    python3 exportar_datos.py --anonimizar
Reemplaza los códigos de acceso por seudónimos (P01, P02, ...) en todas
las tablas, manteniendo la correspondencia entre ellas. Útil para adjuntar
los datos al informe sin exponer identificadores reales.
"""

import os
import sys
import csv
from datetime import datetime
from pathlib import Path

# ── Cargar DATABASE_URL igual que db.py ──────────────────────────────────
try:
    from dotenv import load_dotenv
    load_dotenv()
except ImportError:
    pass  # si no está dotenv, se usa la env var del sistema

try:
    from sqlalchemy import create_engine, text
except ImportError:
    print("ERROR: falta SQLAlchemy. Activá el venv del proyecto:")
    print("  cd ~/chatbot-shi && source venv/bin/activate")
    sys.exit(1)


def _default_sqlite_url():
    db_path = Path(__file__).resolve().parent / "pishico.db"
    return f"sqlite:///{db_path}"


DATABASE_URL = os.getenv("DATABASE_URL", _default_sqlite_url())

# Tablas a exportar, en orden lógico
TABLAS = ["usuarios", "sesiones", "progreso_vocabulario", "progreso_cuento"]

ANONIMIZAR = "--anonimizar" in sys.argv


def exportar():
    print(f"Conectando a: {DATABASE_URL.split('@')[-1] if '@' in DATABASE_URL else DATABASE_URL}")
    engine = create_engine(DATABASE_URL)

    ts = datetime.now().strftime("%Y%m%d_%H%M")
    outdir = Path(f"export_{ts}")
    outdir.mkdir(exist_ok=True)

    datos = {}          # tabla -> (columnas, filas)
    mapa_seudonimos = {}  # codigo_acceso real -> P01, P02...

    with engine.connect() as conn:
        # 1) Leer todas las tablas
        for tabla in TABLAS:
            try:
                res = conn.execute(text(f"SELECT * FROM {tabla}"))
                cols = list(res.keys())
                filas = [list(r) for r in res.fetchall()]
                datos[tabla] = (cols, filas)
                print(f"  {tabla}: {len(filas)} filas")
            except Exception as e:
                print(f"  ⚠️  {tabla}: no se pudo leer ({e})")
                datos[tabla] = ([], [])

        # 2) Construir mapa de seudónimos si se pidió anonimizar
        if ANONIMIZAR:
            cols_u, filas_u = datos.get("usuarios", ([], []))
            if "codigo_acceso" in cols_u:
                idx = cols_u.index("codigo_acceso")
                codigos = sorted({f[idx] for f in filas_u if f[idx]})
                mapa_seudonimos = {c: f"P{i+1:02d}" for i, c in enumerate(codigos)}
                print(f"\nAnonimizando {len(mapa_seudonimos)} participantes...")

    # 3) Aplicar anonimización (codigo_acceso y sender_id)
    if ANONIMIZAR and mapa_seudonimos:
        for tabla, (cols, filas) in datos.items():
            for campo in ("codigo_acceso", "sender_id"):
                if campo in cols:
                    i = cols.index(campo)
                    for f in filas:
                        if f[i] in mapa_seudonimos:
                            f[i] = mapa_seudonimos[f[i]]
        # quitar columnas potencialmente identificatorias de usuarios
        cols_u, filas_u = datos.get("usuarios", ([], []))
        for sensible in ("nombre", "notas_investigador"):
            if sensible in cols_u:
                i = cols_u.index(sensible)
                for f in filas_u:
                    f[i] = ""

    # 4) Escribir CSVs
    for tabla, (cols, filas) in datos.items():
        if not cols:
            continue
        ruta = outdir / f"{tabla}.csv"
        with open(ruta, "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.writer(fh)
            w.writerow(cols)
            w.writerows(filas)
    print(f"\nCSVs escritos en {outdir}/")

    # 5) Escribir Excel con una hoja por tabla
    try:
        import openpyxl
        wb = openpyxl.Workbook()
        wb.remove(wb.active)
        for tabla, (cols, filas) in datos.items():
            if not cols:
                continue
            ws = wb.create_sheet(title=tabla[:31])
            ws.append(cols)
            for f in filas:
                # openpyxl no acepta algunos tipos; convertir a str lo raro
                ws.append([v if isinstance(v, (int, float, str, type(None))) else str(v) for v in f])
            # negrita en encabezado
            for c in ws[1]:
                c.font = openpyxl.styles.Font(bold=True)
            ws.freeze_panes = "A2"
        xlsx = outdir / "pishico_respaldo.xlsx"
        wb.save(xlsx)
        print(f"Excel escrito: {xlsx}")
    except ImportError:
        print("(openpyxl no instalado: se omitió el Excel; los CSV están listos)")

    # 6) Resumen
    lineas = [f"Respaldo Pishico — {datetime.now().strftime('%Y-%m-%d %H:%M')}", ""]
    if ANONIMIZAR:
        lineas.append("MODO ANONIMIZADO: códigos reemplazados por seudónimos (P01, P02, ...)")
        lineas.append("")
    for tabla, (cols, filas) in datos.items():
        lineas.append(f"{tabla}: {len(filas)} registros")
        # rango de fechas si hay columna fecha
        for campo in ("fecha", "fecha_inicio", "fecha_creacion"):
            if campo in cols and filas:
                i = cols.index(campo)
                fechas = [f[i] for f in filas if f[i]]
                if fechas:
                    lineas.append(f"   rango {campo}: {min(fechas)} → {max(fechas)}")
                break
    # participantes distintos
    cols_v, filas_v = datos.get("progreso_vocabulario", ([], []))
    if "sender_id" in cols_v and filas_v:
        i = cols_v.index("sender_id")
        lineas.append(f"\nParticipantes con actividad en vocabulario: {len({f[i] for f in filas_v})}")
    resumen = outdir / "resumen.txt"
    resumen.write_text("\n".join(lineas), encoding="utf-8")
    print(f"Resumen: {resumen}")
    print("\n".join(lineas))

    print(f"\n✅ Listo. Carpeta de respaldo: {outdir.resolve()}")


if __name__ == "__main__":
    exportar()
